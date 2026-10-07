"""menu-belajar.csv → menu-belajar.xlsx (rencana + pelacak + prompt per unit).

Jalankan: python3 -m pip install openpyxl && python3 docs/blueprint/build_xlsx.py
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
rows = list(csv.DictReader(open(f'{ROOT}/docs/blueprint/menu-belajar.csv')))

GRADE_CODE = {'PAUD': 'prek', 'TK': 'tk', 'Kelas 1': 'sd1'}
DOMAIN = {'Baca Tulis': 'literasi', 'Berhitung': 'math', 'English': 'english', 'Sains': 'sains'}
FASE = {'PAUD': 'G4', 'TK': 'G5', 'Kelas 1': 'G6'}
MODELS = [
    (1, 'Dengar lalu pilih gambar', 'pick-one', 'Ada'), (2, 'Ketuk semua', 'tap-all', 'Ada'),
    (3, 'Hitung lalu pilih angka', 'pick-one', 'Ada'), (4, 'Susun benda / kubus', 'build', 'Ada'),
    (5, 'Garis bilangan', 'number-line', 'Ada'), (6, 'Urutkan', 'order', 'Ada'),
    (7, 'Kelompokkan', 'group', 'Ada'), (8, 'Pasangkan', 'match', 'Ada'),
    (9, 'Isian angka', 'number-input', 'Ada'), (10, 'Tebalkan huruf/angka', 'trace', 'Baru'),
    (11, 'Susun huruf jadi kata', 'spell', 'Baru'), (12, 'Gabung suku kata', 'pick-one + animasi', 'Baru (tampilan)'),
    (13, 'Cari di gambar besar', 'scene-tap', 'Baru'), (14, 'Kartu memori', 'memory', 'Baru'),
    (15, 'Yang berbeda', 'pick-one', 'Bisa sekarang (template)'), (16, 'Lihat sekilas', 'pick-one + waktu tampil', 'Baru'),
    (17, 'Timbangan', 'pick-one + visual', 'Baru (visual)'), (18, 'Bagian-keseluruhan', 'pick-one / build', 'Baru (visual)'),
    (19, 'Labirin angka/huruf', 'path', 'Baru'), (20, 'Lengkapi pola', 'pick-one', 'Ada'),
    (21, 'Benar atau tidak', 'pick-one ya/tidak', 'Ada'), (22, 'Baca lalu pilih gambar', 'pick-one', 'Ada'),
    (23, 'Dengar kalimat lalu ketuk', 'scene-tap', 'Baru'), (24, 'Cerita mini', 'order / pick-one', 'Baru (tampilan)'),
]
SCREENS = [
    ('kenalan', 'Huruf/angka besar, Momo menuliskannya (animasi goresan), lalu dibacakan', 'Kartu huruf "a"'),
    ('bunyi', 'Ketuk huruf atau suku kata untuk mendengar bunyinya', 'a, i, u, e, o'),
    ('kata', 'Gambar + kata, suku kata disorot satu per satu saat dibacakan', 'a-pel, a-yam'),
    ('gabung', 'Dua suku kata bergerak lalu menyatu menjadi kata', 'ba + ju = baju'),
    ('cerita', 'Kalimat pendek, kata disorot saat dibacakan; kata bisa diketuk', 'Ini bola. Bola itu merah.'),
    ('coba', 'Mencoba tanpa nilai: tebalkan, ketuk yang dibacakan', 'Tebalkan huruf a'),
    ('ingat', 'Ringkasan satu layar, lalu tombol Ayo latihan', 'Huruf a: apel, ayam, api'),
]
STYLE = ('Ilustrasi datar ramah anak usia 4–7 tahun, garis tepi tebal gelap, warna lembut cerah, bentuk sederhana '
         'dan bulat, satu objek di tengah, latar transparan, tanpa huruf, angka, logo, atau merek, tidak menakutkan.')

def gpt_prompt(r):
    kode, grade, dom = r['kode'], GRADE_CODE[r['grade']], DOMAIN[r['mapel']]
    soal = r['tantangan']
    if soal.startswith('Ada'):
        bagian_soal = (f'BAGIAN B (soal): tidak perlu. Tantangan memakai buku yang sudah ada ({soal[5:]}).')
    else:
        bagian_soal = (
            f'BAGIAN B (soal): buat 10 file skill JSON (Level 1–10, mudah ke sulit, Level 10 tantangan gabungan) '
            f'format Udakids: id "{dom}.{grade}.<kode><n>.<slug>", domain "{dom}", grade "{grade}", '
            f'family "manual" (atau family lain yang sudah ada), tier "basic" untuk PAUD/TK. '
            f'Setiap level minimal 24 soal unik, 3–4 pilihan, satu jawaban benar, pengecoh dari miskonsepsi umum, '
            f'setiap soal punya "say" (dibacakan), "reteach", dan "source". Gunakan hanya interaksi yang sudah ada '
            f'(pick-one, tap-all, order, group, match, build, number-line, number-input).')
    return f"""Kamu penulis konten Udakids (aplikasi belajar anak 4–7 tahun, Indonesia).
UNIT {kode} · {r['grade']} · {r['mapel']} · {r['unit']}
Tujuan: {r['tujuan']}. Rujukan: {r['rujukan']}.
Tema kosakata: {r['tema']}. Model latihan: {r['berlatih']}. Layar pelajaran: {r['belajar']}.

ATURAN WAJIB: bahasa Indonesia sederhana (English: perintah Indonesia, kata target English British); kalimat pendek;
semua teks bisa dibacakan; maksimal 4 pilihan; tanpa kata "salah/gagal"; pujian menyebut usaha; tanpa batas waktu;
tanpa emoji; tidak menyalin soal/gambar dari IXL, Code.org, buku, atau lembar kerja pihak lain; konteks lokal
(nama Indonesia, Rupiah); membaca bahasa Indonesia lewat suku kata; jangan menyimpan atau meminta data anak.

BAGIAN A (pelajaran): JSON {{"kode":"{kode}","judul":"...","layar":[...]}} berisi 3–6 layar. Jenis layar hanya:
kenalan, bunyi, kata, gabung, cerita, coba, ingat. Setiap layar: "jenis", "teks" (≤ 120 huruf), "suara" (kalimat
yang dibacakan), "gambar" (daftar id kata benda dari kamus, huruf kecil, tanpa spasi), dan untuk "kata" isi
"sukuKata" (mis. "a-pel"). Layar terakhir "ingat".

{bagian_soal}

BAGIAN C (gambar): daftar kata benda baru yang butuh gambar, format "id | kata Indonesia | kata Inggris | tema".
Keluarkan hanya JSON untuk A dan B, lalu daftar C."""

wb = Workbook()
thin = Side(style='thin', color='D0D4DC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='3D2C8D')
head_font = Font(bold=True, color='FFFFFF')
wrap = Alignment(wrap_text=True, vertical='top')
MAPEL_FILL = {'Baca Tulis': 'FFF4D6', 'Berhitung': 'E3F2FD', 'English': 'EDE7F6', 'Sains': 'E8F5E9'}

def header(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment, c.border = head_fill, head_font, Alignment(wrap_text=True, vertical='center'), border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'

# --- Petunjuk ---
ws = wb.active
ws.title = 'Petunjuk'
ws.column_dimensions['A'].width = 110
lines = [
    ('Menu Belajar Udakids: PAUD, TK, Kelas 1', True),
    ('Sumber: docs/rencana-gudang-gambar-menu-belajar.md · data: docs/blueprint/menu-belajar.csv · petunjuk lengkap: docs/blueprint/PETUNJUK-GENERATE.md', False),
    ('', False),
    ('CARA MEMAKAI FILE INI', True),
    ('1. Buka sheet "Menu Belajar". Satu baris = satu unit (Belajar → Berlatih → Tantangan → Worksheet).', False),
    ('2. Pakai filter di baris judul untuk memilih Jenjang, Mapel, Fase, atau Status.', False),
    ('3. Generate satu unit: di Claude Code (VS Code), ketik isi kolom "Perintah Claude Code", mis. /generate-unit P-BT-04.', False),
    ('   Alternatif tanpa Claude Code: salin kolom "Prompt GPT" ke ChatGPT, lalu serahkan hasilnya ke tim pengembang untuk divalidasi.', False),
    ('4. Setelah selesai, ubah kolom Status: Belum → Draf → Review → Terbit. Isi Tanggal, PIC, dan Catatan.', False),
    ('5. Sheet "Ringkasan" menghitung progres otomatis dari kolom Status.', False),
    ('', False),
    ('URUTAN YANG DISARANKAN', True),
    ('• Fase G4 = PAUD, G5 = TK, G6 = Kelas 1. Kerjakan berurutan dari kode terkecil (unit awal jadi dasar unit berikutnya).', False),
    ('• Soal "BARU 10 level" bisa digenerate sekarang. Pelajaran (Belajar) disimpan sebagai draf sampai pemutar pelajaran selesai dibangun (fase G3).', False),
    ('• Model latihan berstatus "Baru" (sheet Model Soal) belum ada di aplikasi; sementara soal memakai model yang sudah ada.', False),
    ('', False),
    ('ARTI KOLOM', True),
    ('Belajar (layar): jenis layar pelajaran (sheet Layar Pelajaran). Berlatih (model): nomor di sheet Model Soal.', False),
    ('Tantangan: "Ada: <buku> <kode>" = memakai soal yang sudah ada; "BARU 10 level" = topik soal baru.', False),
    ('Worksheet: jenis lembar kerja cetak (PDF) yang nanti diunduh orang tua.', False),
    ('', False),
    ('ATURAN YANG TIDAK BOLEH DILANGGAR', True),
    ('Tanpa data pribadi anak · tanpa batas waktu/nyawa/streak · tanpa kata "salah/gagal" · maks 4 pilihan · semua teks dibacakan · '
     'tidak menyalin soal/gambar pihak lain · tanpa AI di area anak · tanpa harga di area anak · semua hasil direview sebelum terbit.', False),
]
for text, bold in lines:
    ws.append([text])
    c = ws.cell(row=ws.max_row, column=1)
    c.alignment = wrap
    if bold:
        c.font = Font(bold=True, size=13 if ws.max_row == 1 else 11, color='3D2C8D')

# --- Menu Belajar ---
ws = wb.create_sheet('Menu Belajar')
cols = ['Kode', 'Jenjang', 'Mapel', 'Unit', 'Tujuan', 'Rujukan', 'Belajar (layar)', 'Berlatih (model)', 'Tantangan (soal)',
        'Tema kosakata', 'Worksheet cetak', 'Pekerjaan', 'Fase', 'Status', 'Tanggal', 'PIC', 'Catatan',
        'Perintah Claude Code', 'Prompt GPT']
widths = [10, 9, 11, 30, 38, 32, 22, 12, 40, 24, 20, 26, 7, 10, 12, 12, 24, 26, 60]
header(ws, cols, widths)
for r in rows:
    ws.append([r['kode'], r['grade'], r['mapel'], r['unit'], r['tujuan'], r['rujukan'], r['belajar'], r['berlatih'],
               r['tantangan'], r['tema'], r['cetak'], r['status'], FASE[r['grade']], 'Belum', None, None, None,
               f"/generate-unit {r['kode']}", gpt_prompt(r)])
    i = ws.max_row
    fill = PatternFill('solid', fgColor=MAPEL_FILL[r['mapel']])
    for j in range(1, len(cols) + 1):
        c = ws.cell(row=i, column=j)
        c.alignment, c.border = wrap, border
        if j <= 4:
            c.fill = fill
    ws.row_dimensions[i].height = 60
    ws.cell(row=i, column=15).number_format = 'yyyy-mm-dd'
last = ws.max_row
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{last}'
dv = DataValidation(type='list', formula1='"Belum,Draf,Review,Terbit,Ditunda"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f'N2:N{last}')
for val, color in [('Draf', 'FFF59D'), ('Review', 'FFCC80'), ('Terbit', 'A5D6A7'), ('Ditunda', 'E0E0E0')]:
    ws.conditional_formatting.add(f'N2:N{last}', CellIsRule(operator='equal', formula=[f'"{val}"'],
                                                             fill=PatternFill('solid', fgColor=color)))

# --- Ringkasan ---
ws = wb.create_sheet('Ringkasan')
header(ws, ['Jenjang', 'Mapel', 'Unit', 'Soal baru', 'Draf', 'Review', 'Terbit', 'Progres'], [10, 12, 8, 10, 8, 8, 8, 10])
R = "'Menu Belajar'"
for g in ['PAUD', 'TK', 'Kelas 1']:
    for m in ['Baca Tulis', 'Berhitung', 'English', 'Sains']:
        n = ws.max_row + 1
        crit = f'{R}!$B:$B,"{g}",{R}!$C:$C,"{m}"'
        ws.append([g, m, f'=COUNTIFS({crit})', f'=COUNTIFS({crit},{R}!$I:$I,"BARU*")',
                   f'=COUNTIFS({crit},{R}!$N:$N,"Draf")', f'=COUNTIFS({crit},{R}!$N:$N,"Review")',
                   f'=COUNTIFS({crit},{R}!$N:$N,"Terbit")', f'=IF(C{n}=0,0,G{n}/C{n})'])
        ws.cell(row=n, column=8).number_format = '0%'
n = ws.max_row + 1
ws.append(['Total', '', f'=SUM(C2:C{n-1})', f'=SUM(D2:D{n-1})', f'=SUM(E2:E{n-1})', f'=SUM(F2:F{n-1})',
           f'=SUM(G2:G{n-1})', f'=IF(C{n}=0,0,G{n}/C{n})'])
ws.cell(row=n, column=8).number_format = '0%'
for c in ws[n]:
    c.font = Font(bold=True)

# --- Model Soal ---
ws = wb.create_sheet('Model Soal')
header(ws, ['No', 'Model', 'Interaksi', 'Status di aplikasi'], [6, 30, 26, 26])
for m in MODELS:
    ws.append(list(m))

# --- Layar Pelajaran ---
ws = wb.create_sheet('Layar Pelajaran')
header(ws, ['Jenis layar', 'Isi', 'Contoh'], [14, 70, 30])
for s in SCREENS:
    ws.append(list(s))
    for c in ws[ws.max_row]:
        c.alignment = wrap

# --- Prompt Gambar ---
ws = wb.create_sheet('Prompt Gambar')
header(ws, ['Bagian', 'Isi'], [22, 110])
for k, v in [
    ('Gaya Udakids v1', STYLE),
    ('Template satu gambar', '{GAYA} Objek: {kata Indonesia} ({kata Inggris}), {varian}.'),
    ('Template grid 2×2 (hemat 75%)', '{GAYA} Empat objek terpisah dalam grid 2×2 dengan jarak putih lebar, masing-masing di tengah selnya: '
     '1) {kata1}, 2) {kata2}, 3) {kata3}, 4) {kata4}. Semua objek ukuran serupa dan tidak saling menyentuh.'),
    ('Contoh varian', 'apel: merah utuh · hijau utuh · apel terpotong · apel di piring'),
    ('Pengaturan API', 'model gpt-image-1-mini · quality medium (pratinjau: low) · size 1024x1024 · background transparent · '
     'output_format webp · lewat Batch API'),
    ('Setelah generate', 'potong grid, ubah ke WebP 512 dan 256, review admin, catat lisensi dan prompt di Kamus Bergambar'),
]:
    ws.append([k, v])
    for c in ws[ws.max_row]:
        c.alignment = wrap

wb.save(f'{ROOT}/docs/blueprint/menu-belajar.xlsx')
print('ok', len(rows))
